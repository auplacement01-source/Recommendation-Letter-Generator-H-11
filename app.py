from __future__ import annotations

import hmac
import io
import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from supabase import Client, create_client


# ------------------------------ App configuration ------------------------------
st.set_page_config(
    page_title="Air University | Letter Desk",
    page_icon="✉️",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    :root { --ink:#172b4d; --blue:#0c5c91; --pale:#eff6fb; --line:#dce5ed; }
    .stApp { background: linear-gradient(180deg,#f8fbfd 0%,#ffffff 360px); }
    [data-testid="stSidebar"] { background:#f1f6fa; border-right:1px solid #e2eaf0; }
    .hero { padding:1.35rem 1.55rem; border:1px solid #dce8f1; border-radius:18px;
            background:linear-gradient(120deg,#ffffff 0%,#eef6fb 100%); margin-bottom:1rem; }
    .hero h1 { color:#163b5b; margin:0 0 .3rem 0; font-size:2rem; }
    .hero p { color:#52677a; margin:0; }
    .subtle { color:#63788a; font-size:.92rem; }
    .step-card { padding:.8rem 1rem; border:1px solid #e0e9f0; border-radius:12px; background:#fff; }
    div[data-testid="stMetric"] { background:#fff; border:1px solid #e3ebf1; border-radius:12px; padding:.65rem .85rem; }
    .stButton > button[kind="primary"] { background:#0c5c91; border-color:#0c5c91; }
    </style>
    """,
    unsafe_allow_html=True,
)

PROGRAMS: dict[str, str] = {
    "Master of Science in Clinical Psychology": "MSCP",
    "Bachelor of Business Administration": "BBA",
    "Bachelor of Science in Psychology": "BSP",
    "Master of Business Administration": "MBA",
    "Bachelor of Science in Computer Science": "BSCS",
    "Other / enter below": "OTHER",
}
SEMESTERS = list(range(1, 13))
DEPARTMENT_PRESETS = [
    "IBD/AU/DSA/PLAC/H-11",
    "IBD/AU/PLAC/H-11",
    "Custom prefix",
]

# This intentionally small, editable offline dictionary makes a best-effort guess.
# Names can be unisex or differently gendered by family; always allow a human override.
FEMALE_NAMES = set("""
Ayesha Fatima Zainab Maryam Mariam Amina Aaminah Amna Sana Hira Iqra Mahnoor Noor
Zubda Zubda Fatima Kazam Imaan Iman Amnah Amna Siddiqui Laiba Eman Emaan Anum
Anam Hania Haniya Dua Aleena Alina Aliza Areeba Arooba Arooj Mehak Mahak Saira
Sadia Nadia Rabia Rabiya Sidra Saba Kiran Kiran Bushra Beenish Sehrish Zoya Zara
Hafsa Hafsa Anaya Anayah Minahil Minal Nimra Nimrah Momina Mominah Muskan Rida
Ridaa Sehr Hoorain Aiman Ayra Aiza Aqsa Aqsa Rameen Rimsha Saniya Sania Komal
Maham Maira Misha Huda Hiba Hifza Esha Ayesha Anisha Areej Erum Irum Sumbal
Shazia Farah Farwa Fiza Fariha Faryal Arisha Hareem Hareem Alishba Ayesha Noor
Aqsa Javeria Jaweria Khadija Khadijah Aaliyah Aliya Sehrish Muneeba Muneebah
Nawal Nida Noreen Naila Naima Naeema Samina Samra Samreen Shumaila Sumaira
Tahira Tehreem Tuba Tooba Urooj Yasmin Yusra Zunaira Zunairah Zainab Zahra
Abeer Arooj Afsheen Afreen Alishba Aqila Bushra Faiza Fiza Gul Hina Iffat
Jannat Kanwal Lubna Mahwish Mehwish Nargis Parisa Rabail Rukhsana Saman Shabana
Shagufta Sharmeen Shazia Tehmina Uzma Wajiha Warda Wania Yumna
""".split())
MALE_NAMES = set("""
Muhammad Mohammad Mohammed Ahmed Ahmad Ali Hassan Hussain Husain Abdullah Abdul
Abdulrahman Usman Uthman Umar Omar Bilal Hamza Ibrahim Ismail Ismail Imran Imran
Ahsan Asad Salman Zain Zayn Saad Shahzaib Shehryar Talha Taha Tayyab Tayab Yasir
Yasir Farhan Fahad Faizan Faisal Fahim Fahad Muneeb Muneeb Daniyal Danish Dawood
Dawood Awais Owais Arslan Arsalan Haris Harris Huzaifa Huzayfa Junaid Junayd
Kamran Kashif Khurram Luqman Moiz Moeed Mustafa Musab Noman Nouman Qasim Qasim
Raza Rehan Rizwan Saif Sameer Samir Shahbaz Shahid Sohaib Suhaib Usama Waleed
Walid Yousuf Yusuf Zubair Zain Ali Abbas Adnan Adeel Aftab Akbar Amir Ameer
Anas Aqib Arham Asif Atif Azlan Azhar Babar Basit Burhan Ehtesham Ehsan
Faisal Ghazi Ghulam Haider Hamid Haroon Hashim Hammad Haseeb Ilyas Irfan Jibran
Khalid Majid Malik Mansoor Masood Mehmood Nabeel Nabil Naveed Nawaz Niaz
Parvez Pervez Raheel Rashid Rauf Shayan Sohail Suhail Tariq Waqar Waqas
""".split())

LOGO_PATH = Path(__file__).with_name("air_university_logo.png")


def secret_value(section: str, key: str, env_name: str | None = None) -> str:
    """Read a value from Streamlit secrets, falling back to an environment variable."""
    try:
        section_data = st.secrets.get(section, {})
        value = section_data.get(key, "") if hasattr(section_data, "get") else ""
        if value:
            return str(value).strip()
    except Exception:
        pass
    if env_name:
        return os.getenv(env_name, "").strip()
    return ""


@st.cache_resource(show_spinner=False)
def get_supabase(url: str, key: str) -> Client:
    return create_client(url, key)


def require_password() -> None:
    configured = secret_value("app", "password", "APP_PASSWORD")
    if not configured:
        st.error("**Setup required:** add a strong shared password as `APP_PASSWORD` in Streamlit Secrets before using this app.")
        st.code('[app]\npassword = "replace-with-a-long-random-password"', language="toml")
        st.stop()
    if st.session_state.get("authenticated"):
        return
    st.markdown(
        '<div class="hero"><h1>Air University · Letter Desk</h1>'
        '<p>Secure recommendation and internship letter workspace</p></div>',
        unsafe_allow_html=True,
    )
    with st.form("login_form"):
        entered = st.text_input("Office password", type="password")
        submitted = st.form_submit_button("Unlock workspace", type="primary", use_container_width=True)
    if submitted:
        if hmac.compare_digest(entered.encode("utf-8"), configured.encode("utf-8")):
            st.session_state["authenticated"] = True
            st.rerun()
        else:
            st.error("That password does not match. Please try again.")
    st.caption("This is a simple shared-password gate for a small office workflow. Keep the app private and do not share the Supabase secret key.")
    st.stop()


def get_client_or_stop() -> Client:
    url = secret_value("supabase", "url", "SUPABASE_URL")
    key = (
        secret_value("supabase", "key", "SUPABASE_SECRET_KEY")
        or secret_value("supabase", "service_role_key", "SUPABASE_SERVICE_ROLE_KEY")
    )
    if not url or not key:
        st.error("Supabase is not configured yet. Add your project URL and server-side secret key in Streamlit Secrets, then restart the app.")
        st.code('[supabase]\nurl = "https://YOUR-PROJECT-REF.supabase.co"\nkey = "YOUR-SUPABASE-SERVER-ONLY-SECRET-KEY"', language="toml")
        st.info("Do not put the key in GitHub, in app.py, or in a browser-side script. Follow the README setup steps.")
        st.stop()
    try:
        return get_supabase(url, key)
    except Exception as exc:
        st.error(f"Could not connect to Supabase: {exc}")
        st.stop()


def ordinal(n: int) -> str:
    if 10 <= n % 100 <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def infer_gender(full_name: str) -> tuple[str, str]:
    """Return a best-effort value from the first name; unknown names remain neutral."""
    tokens = re.findall(r"[A-Za-zÀ-ÿ'-]+", (full_name or "").strip())
    if not tokens:
        return "They", "No name to analyze"
    # Pakistani full names may start with a given name, title, or compound name.
    first = tokens[0].strip("'-").title()
    normalized = first.replace("'", "")
    if normalized in FEMALE_NAMES:
        return "She", "Best-effort name match"
    if normalized in MALE_NAMES:
        return "He", "Best-effort name match"
    if len(tokens) > 1:
        second = tokens[1].strip("'-").title().replace("'", "")
        if second in FEMALE_NAMES:
            return "She", "Best-effort match on second name"
        if second in MALE_NAMES:
            return "He", "Best-effort match on second name"
    return "They", "Name not recognized—please choose pronouns"


def resolve_gender(name: str, explicit: Any = None) -> tuple[str, str]:
    options = {"female": "She", "she": "She", "f": "She", "male": "He", "he": "He", "m": "He", "they": "They", "other": "They"}
    if explicit is not None and not pd.isna(explicit) and str(explicit).strip():
        given = str(explicit).strip().lower()
        if given in options:
            return options[given], "Excel override"
        # Also accept a full selector label, e.g. "Female (She/Her)".
        if "she" in given or "female" in given:
            return "She", "Excel override"
        if "he" in given or "male" in given:
            return "He", "Excel override"
        if "they" in given:
            return "They", "Excel override"
    return infer_gender(name)


def safe_text(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def make_row(
    *,
    student_name: str,
    registration_id: str,
    gender: str,
    program: str,
    program_code: str,
    semester: int,
    department_prefix: str,
    letter_date: date,
    recipient_mode: str,
    company_name: str = "",
    city: str = "",
) -> dict[str, Any]:
    return {
        "student_name": student_name.strip(),
        "registration_id": registration_id.strip(),
        "gender": gender,
        "program": program,
        "program_code": program_code,
        "semester": int(semester),
        "department_prefix": department_prefix.strip().strip("/"),
        "letter_date": letter_date.isoformat(),
        "recipient_mode": recipient_mode,
        "company_name": company_name.strip() if recipient_mode == "Company / organization" else "",
        "city": city.strip() if recipient_mode == "Company / organization" else "",
    }


def allocate_and_save(client: Client, rows: list[dict[str, Any]], include_semester: bool) -> list[dict[str, Any]]:
    """Allocate unique numbers atomically in Postgres, then save each letter record."""
    saved: list[dict[str, Any]] = []
    for row in rows:
        yr = date.fromisoformat(row["letter_date"]).strftime("%y")
        try:
            rpc_result = client.rpc(
                "next_letter_serial",
                {
                    "p_department_prefix": row["department_prefix"],
                    "p_program_code": row["program_code"],
                    "p_year_2d": yr,
                },
            ).execute()
            sequence = rpc_result.data
            if isinstance(sequence, list):
                sequence = sequence[0]
            sequence = int(sequence)
            parts = [row["department_prefix"], row["program_code"], yr]
            if include_semester:
                parts.append(f"S{row['semester']}")
            parts.append(f"{sequence:03d}")
            record = {
                **row,
                "year_2d": yr,
                "sequence_no": sequence,
                "letter_number": "/".join(parts),
            }
            result = client.table("letters").insert(record).execute()
            if not result.data:
                raise RuntimeError("Supabase did not return the inserted letter record.")
            saved.append(result.data[0])
        except Exception as exc:
            raise RuntimeError(
                f"Could not save {row.get('student_name', 'student')} ({row.get('registration_id', '')}): {exc}"
            ) from exc
    return saved


def format_letter_date(value: Any) -> str:
    if isinstance(value, date):
        d = value
    else:
        d = date.fromisoformat(str(value)[:10])
    return f"{ordinal(d.day)} {d.strftime('%B, %Y')}"


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        from docx.oxml import OxmlElement
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(25, 35, 45)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.12


def add_letter(doc: Document, record: dict[str, Any], final_page: bool = False) -> None:
    header = doc.add_table(rows=1, cols=2)
    header.autofit = False
    header.columns[0].width = Inches(4.55)
    header.columns[1].width = Inches(2.0)
    remove_table_borders(header)
    left, right = header.rows[0].cells
    left.width = Inches(4.55)
    right.width = Inches(2.0)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(left, 0, 0, 0, 0)
    set_cell_margins(right, 0, 0, 0, 0)
    left_p = left.paragraphs[0]
    left_p.paragraph_format.space_after = Pt(2)
    if LOGO_PATH.exists():
        left_p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.0))
    else:
        run = left_p.add_run("AIR UNIVERSITY")
        run.bold = True
        run.font.color.rgb = RGBColor(0, 91, 145)
        run.font.size = Pt(12)
    number_p = left.add_paragraph()
    number_p.paragraph_format.space_before = Pt(0)
    number_p.paragraph_format.space_after = Pt(0)
    nr = number_p.add_run(safe_text(record.get("letter_number")))
    nr.font.size = Pt(8.5)
    date_p = right.paragraphs[0]
    date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_p.paragraph_format.space_after = Pt(0)
    dr = date_p.add_run(format_letter_date(record.get("letter_date", date.today().isoformat())))
    dr.font.size = Pt(9.5)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(5)
    mode = safe_text(record.get("recipient_mode"))
    company = safe_text(record.get("company_name"))
    city = safe_text(record.get("city"))
    if mode == "Company / organization" and company:
        to = doc.add_paragraph()
        to.paragraph_format.space_after = Pt(0)
        r = to.add_run("HR Manager")
        r.bold = True
        org = doc.add_paragraph()
        org.paragraph_format.space_after = Pt(0)
        org.add_run(company)
        if city:
            loc = doc.add_paragraph()
            loc.paragraph_format.space_after = Pt(0)
            loc.add_run(city)
    salutation = doc.add_paragraph()
    salutation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    salutation.paragraph_format.space_before = Pt(8)
    salutation.paragraph_format.space_after = Pt(15)
    s = salutation.add_run("To Whom it May Concern")
    s.bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(20)
    t = title.add_run("REQUEST FOR INTERNSHIP OPPORTUNITY FOR AIR UNIVERSITY STUDENT")
    t.bold = True
    t.underline = True
    t.font.size = Pt(10.5)

    name = safe_text(record.get("student_name"))
    reg = safe_text(record.get("registration_id"))
    gender = safe_text(record.get("gender")) or "They"
    # Sentence pronoun in the source template is subject case: He / She (or neutral They).
    program = safe_text(record.get("program"))
    semester = int(record.get("semester") or 1)
    body = doc.add_paragraph()
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bold_prefix = body.add_run(f"{name}, Registration No. {reg},")
    bold_prefix.bold = True
    body.add_run(
        f" is a student at Air University, a federally chartered public sector university, recognized by HEC. "
        f"{gender} is currently enrolled in the {ordinal(semester)} Semester of {program}."
    )

    organization_phrase = f"at {company}" if mode == "Company / organization" and company else "at your esteemed organization"
    paragraph = doc.add_paragraph(
        f"It is requested that the student may please be provided with an internship opportunity {organization_phrase}. "
        "The student will benefit a great deal by working at an organization with a good reputation, which would go a long way in professional development."
    )
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph("If you have any internship/ job opportunity in future, please share with us.")

    sign_table = doc.add_table(rows=1, cols=2)
    sign_table.autofit = False
    sign_table.columns[0].width = Inches(3.65)
    sign_table.columns[1].width = Inches(2.9)
    remove_table_borders(sign_table)
    blank, sign = sign_table.rows[0].cells
    blank.width = Inches(3.65)
    sign.width = Inches(2.9)
    for cell in (blank, sign):
        set_cell_margins(cell, 0, 0, 0, 0)
    sign_p = sign.paragraphs[0]
    sign_p.paragraph_format.space_before = Pt(78)
    sign_p.paragraph_format.space_after = Pt(0)
    sign_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    officer = safe_text(record.get("officer_name")) or "MOBEEN JAMSHED"
    role = safe_text(record.get("officer_title")) or "Assistant Director"
    office = safe_text(record.get("office_name")) or "Placement & Alumni Affairs"
    rr = sign_p.add_run(f"({officer})")
    rr.bold = True
    sign_p.add_run(f"\n{role}\n{office}")


def build_docx(records: list[dict[str, Any]]) -> bytes:
    doc = Document()
    style_document(doc)
    for index, record in enumerate(records):
        if index:
            doc.add_page_break()
        add_letter(doc, record)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def get_year(date_text: str) -> str:
    return date.fromisoformat(date_text).strftime("%y")


def normalize_program(raw: Any, default_program: str) -> tuple[str, str]:
    text = safe_text(raw)
    if not text:
        text = default_program
    for label, code in PROGRAMS.items():
        if text.lower() == label.lower() or text.upper() == code:
            if label == "Other / enter below":
                return "Other / enter below", code
            return label, code
    # For a custom course name, create a stable short code from its words.
    code = re.sub(r"[^A-Z0-9]", "", "".join(part[:1] for part in re.findall(r"[A-Za-z0-9]+", text.upper())))[:8]
    return text, code or "OTHER"


def normalize_semester(raw: Any, default_semester: int) -> int:
    text = safe_text(raw)
    if not text:
        return int(default_semester)
    match = re.search(r"\d+", text)
    if not match:
        raise ValueError(f"Invalid semester value: {text}")
    value = int(match.group())
    if not 1 <= value <= 12:
        raise ValueError(f"Semester must be between 1 and 12: {text}")
    return value


def template_excel() -> bytes:
    example = pd.DataFrame(columns=[
        "student_name", "registration_id", "program", "semester", "gender",
        "company_name", "city", "recipient_mode",
    ])
    info = pd.DataFrame({
        "Column": ["student_name", "registration_id", "program", "semester", "gender", "company_name", "city", "recipient_mode"],
        "Required?": ["Yes", "Yes", "No (uses selected default)", "No (uses selected default)", "No (auto-detects from name)", "No", "No", "No (uses selected default)"],
        "Example / accepted values": ["Ayesha Khan", "2504332", "Master of Science in Clinical Psychology or MSCP", "3 or 3rd Semester", "Female, Male, or They; blank = best-effort", "PIMS", "Islamabad", "Company / organization or General"],
    })
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        example.to_excel(writer, index=False, sheet_name="Students")
        info.to_excel(writer, index=False, sheet_name="Read me")
    return out.getvalue()


def excel_bytes(records: list[dict[str, Any]]) -> bytes:
    df = pd.DataFrame(records)
    preferred = [
        "letter_number", "letter_date", "student_name", "registration_id", "gender",
        "program", "program_code", "semester", "recipient_mode", "company_name", "city",
        "department_prefix", "year_2d", "sequence_no", "created_at",
    ]
    columns = [c for c in preferred if c in df.columns] + [c for c in df.columns if c not in preferred]
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df[columns].to_excel(writer, index=False, sheet_name="Letter records")
        ws = writer.sheets["Letter records"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column_cells in ws.columns:
            max_length = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 45)
            ws.column_dimensions[column_cells[0].column_letter].width = max(12, max_length)
    return out.getvalue()


def fetch_all_records(client: Client, limit: int = 1000) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    start = 0
    while True:
        batch = client.table("letters").select("*").order("created_at", desc=True).range(start, start + limit - 1).execute().data or []
        rows.extend(batch)
        if len(batch) < limit:
            break
        start += limit
    return rows


def date_from_record(record: dict[str, Any]) -> str:
    value = record.get("letter_date")
    if isinstance(value, str):
        return value[:10]
    return date.today().isoformat()


def app_main() -> None:
    require_password()
    client = get_client_or_stop()

    st.markdown(
        '<div class="hero"><h1>Air University · Letter Desk</h1>'
        '<p>Create, store, find and reprint internship recommendation letters — one at a time or from an Excel roster.</p></div>',
        unsafe_allow_html=True,
    )

    with st.sidebar:
        st.subheader("Letter settings")
        prefix_choice = st.selectbox("Department / office prefix", DEPARTMENT_PRESETS, index=0)
        if prefix_choice == "Custom prefix":
            department_prefix = st.text_input("Enter department prefix", value="IBD/AU/DSA/PLAC/H-11")
        else:
            department_prefix = prefix_choice
        include_semester = st.checkbox("Include semester in letter number", value=False, help="Off matches your example: PREFIX/PROGRAM/YY/001")
        st.caption("Example: `IBD/AU/DSA/PLAC/H-11/BBA/26/001`. The year follows the letter date; serials increase separately for each office, program, and year.")
        with st.expander("Signature block", expanded=False):
            officer_name = st.text_input("Officer name", value="MOBEEN JAMSHED")
            officer_title = st.text_input("Officer title", value="Assistant Director")
            office_name = st.text_input("Office", value="Placement & Alumni Affairs")
        st.divider()
        if st.button("Lock workspace", use_container_width=True):
            st.session_state["authenticated"] = False
            st.rerun()
        st.caption("Tip: Keep this workspace password-protected and limit access to authorized office staff.")

    common = {
        "department_prefix": department_prefix,
        "officer_name": officer_name,
        "officer_title": officer_title,
        "office_name": office_name,
    }
    tabs = st.tabs(["Create one", "Bulk from Excel", "Find / reprint", "Records & export"])

    # ------------------------------ Single letter ------------------------------
    with tabs[0]:
        st.markdown("### Create an individual letter")
        st.caption("Enter the student details. Name-based pronoun matching is an offline best-effort guess; review the pronoun selector before generating.")
        col1, col2 = st.columns(2)
        with col1:
            student_name = st.text_input(
                "Student full name *",
                placeholder="e.g. Imaan Shah",
                key="single_name",
                on_change=lambda: st.session_state.update(
                    single_gender=infer_gender(st.session_state.get("single_name", ""))[0]
                ),
            )
            registration_id = st.text_input("Registration ID *", placeholder="e.g. 2504269", key="single_reg")
            program_label = st.selectbox("Program *", list(PROGRAMS.keys()), index=0, key="single_program")
            custom_program = ""
            if program_label == "Other / enter below":
                custom_program = st.text_input("Program name", placeholder="Enter full program name")
            semester = st.selectbox("Semester *", SEMESTERS, index=2, format_func=lambda n: f"{ordinal(n)} Semester", key="single_semester")
        with col2:
            detected_gender, gender_reason = infer_gender(student_name)
            pronoun_choice = st.selectbox(
                "Pronoun for the letter *",
                ["She", "He", "They"],
                index=["She", "He", "They"].index(detected_gender),
                format_func=lambda v: {"She": "She / her", "He": "He / him", "They": "They / them"}[v],
                key="single_gender",
            )
            if student_name:
                if detected_gender == "They":
                    st.warning("We could not confidently match this name. Please select the correct pronoun above.")
                else:
                    st.info(f"Suggested from the name: **{detected_gender}** · {gender_reason}. You can change it.")
            recipient_mode = st.radio("Letter recipient", ["General (any organization)", "Company / organization"], horizontal=True, key="single_recipient")
            company_name = ""
            city = ""
            if recipient_mode == "Company / organization":
                company_name = st.text_input("Company / organization name *", placeholder="e.g. PIMS")
                city = st.text_input("City (optional)", placeholder="e.g. Islamabad")
            letter_date = st.date_input("Letter date", value=date.today(), format="DD/MM/YYYY", key="single_date")
        single_program = custom_program.strip() if program_label == "Other / enter below" else program_label
        if program_label == "Other / enter below":
            _, single_program_code = normalize_program(single_program, single_program)
        else:
            single_program_code = PROGRAMS[program_label]
        generate_one = st.button("Generate and save letter", type="primary", use_container_width=True, key="generate_single")
        if generate_one:
            issues = []
            if not student_name.strip(): issues.append("Enter the student's full name.")
            if not registration_id.strip(): issues.append("Enter the registration ID.")
            if not department_prefix.strip(): issues.append("Enter a department prefix in the sidebar.")
            if program_label == "Other / enter below" and not single_program: issues.append("Enter the custom program name.")
            if recipient_mode == "Company / organization" and not company_name.strip(): issues.append("Enter the company or organization name, or choose General.")
            if issues:
                for item in issues: st.error(item)
            else:
                row = make_row(
                    student_name=student_name, registration_id=registration_id, gender=pronoun_choice,
                    program=single_program, program_code=single_program_code, semester=semester,
                    department_prefix=department_prefix, letter_date=letter_date,
                    recipient_mode=recipient_mode, company_name=company_name, city=city,
                )
                try:
                    with st.spinner("Allocating the next letter number and saving to Supabase…"):
                        saved = allocate_and_save(client, [{**row, **common}], include_semester)[0]
                    st.session_state["last_single_record"] = saved
                    st.success(f"Saved as **{saved['letter_number']}**. You can print it now or find it later by registration ID.")
                except Exception as exc:
                    st.error(str(exc))
        last_single = st.session_state.get("last_single_record")
        if last_single:
            st.download_button(
                "Download / print latest letter (.docx)", build_docx([last_single]),
                file_name=f"{last_single['letter_number'].replace('/', '-')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary",
                key=f"download_single_{last_single['id']}",
            )
            with st.expander("Latest saved record details"):
                st.json({k: last_single.get(k) for k in ("letter_number", "student_name", "registration_id", "gender", "program", "semester", "company_name", "city")})

    # ------------------------------ Bulk workbook ------------------------------
    with tabs[1]:
        st.markdown("### Generate a batch from an Excel roster")
        st.write("Upload `.xlsx` or `.csv`. Each student gets a separate letter number; all letters are saved in Supabase and combined into one downloadable Word document.")
        st.download_button("Download blank Excel template", template_excel(), "student_letter_template.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="template_download")
        bulk1, bulk2, bulk3 = st.columns(3)
        with bulk1:
            default_bulk_program = st.selectbox(
                "Default program for blank cells",
                [program for program in PROGRAMS if program != "Other / enter below"],
                index=0,
                key="bulk_program",
            )
        with bulk2:
            default_bulk_semester = st.selectbox("Default semester for blank cells", SEMESTERS, index=2, format_func=lambda n: f"{ordinal(n)} Semester", key="bulk_semester")
        with bulk3:
            bulk_default_mode = st.selectbox("Default recipient mode", ["General (any organization)", "Company / organization"], key="bulk_mode")
        bulk_date = st.date_input("Date for this batch", value=date.today(), format="DD/MM/YYYY", key="bulk_date")
        uploaded = st.file_uploader("Choose student roster", type=["xlsx", "csv"], key="bulk_upload")
        if uploaded:
            try:
                if uploaded.name.lower().endswith(".csv"):
                    df = pd.read_csv(uploaded, dtype=str).fillna("")
                else:
                    df = pd.read_excel(uploaded, dtype=str).fillna("")
                df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]
                st.caption(f"Loaded {len(df)} row(s). Required columns: `student_name` and `registration_id`.")
                st.dataframe(df.head(15), use_container_width=True, hide_index=True)
                if not {"student_name", "registration_id"}.issubset(df.columns):
                    st.error("The spreadsheet needs columns named `student_name` and `registration_id`. Download the blank template above.")
                else:
                    if st.button("Review and generate all letters", type="primary", key="bulk_generate"):
                        default_program, default_code = normalize_program(default_bulk_program, default_bulk_program)
                        prepared: list[dict[str, Any]] = []
                        problems: list[str] = []
                        seen: set[tuple[str, str, str]] = set()
                        for idx, rowdata in df.iterrows():
                            excel_row = idx + 2
                            name = safe_text(rowdata.get("student_name"))
                            reg = safe_text(rowdata.get("registration_id"))
                            if not name or not reg:
                                problems.append(f"Excel row {excel_row}: missing student name or registration ID; skipped.")
                                continue
                            key = (reg.casefold(), safe_text(rowdata.get("company_name")).casefold(), safe_text(rowdata.get("program")).casefold())
                            if key in seen:
                                problems.append(f"Excel row {excel_row}: duplicate registration/company/program in this upload; skipped.")
                                continue
                            seen.add(key)
                            try:
                                program, code = normalize_program(rowdata.get("program"), default_program)
                                if program == "Other / enter below":
                                    problems.append(f"Excel row {excel_row}: the 'Other' program option needs a full program name in the spreadsheet; skipped.")
                                    continue
                                sem = normalize_semester(rowdata.get("semester"), default_bulk_semester)
                                gender, _ = resolve_gender(name, rowdata.get("gender"))
                                mode_raw = safe_text(rowdata.get("recipient_mode")).lower()
                                if mode_raw in {"company", "organization", "company / organization", "specific"}:
                                    mode = "Company / organization"
                                elif mode_raw in {"general", "general (any organization)", "any"}:
                                    mode = "General (any organization)"
                                else:
                                    mode = bulk_default_mode
                                company = safe_text(rowdata.get("company_name"))
                                city_value = safe_text(rowdata.get("city"))
                                if mode == "Company / organization" and not company:
                                    problems.append(f"Excel row {excel_row}: company mode was selected but company_name is blank; skipped.")
                                    continue
                                prepared.append(make_row(
                                    student_name=name, registration_id=reg, gender=gender,
                                    program=program, program_code=code, semester=sem,
                                    department_prefix=department_prefix, letter_date=bulk_date,
                                    recipient_mode=mode, company_name=company, city=city_value,
                                ))
                            except Exception as exc:
                                problems.append(f"Excel row {excel_row}: {exc}; skipped.")
                        if problems:
                            with st.expander(f"Import notes ({len(problems)})", expanded=True):
                                for problem in problems: st.write(f"- {problem}")
                        if not prepared:
                            st.error("No valid student rows to generate.")
                        else:
                            st.info(f"Ready to generate **{len(prepared)}** letter(s). Rows with issues above will not be processed.")
                            try:
                                with st.spinner(f"Saving {len(prepared)} letters to Supabase…"):
                                    saved_batch = allocate_and_save(client, [{**row, **common} for row in prepared], include_semester)
                                st.session_state["last_bulk_records"] = saved_batch
                                st.success(f"Successfully saved {len(saved_batch)} letter(s). The combined document is ready below.")
                            except Exception as exc:
                                st.error(f"Batch stopped after any earlier successful inserts. Check Records & export before retrying to avoid duplicate letters. Details: {exc}")
                    last_batch = st.session_state.get("last_bulk_records", [])
                    if last_batch:
                        st.dataframe(pd.DataFrame(last_batch)[[c for c in ["letter_number", "student_name", "registration_id", "gender", "program", "semester", "company_name"] if c in last_batch[0]]], use_container_width=True, hide_index=True)
                        combined_bytes = build_docx(last_batch)
                        st.download_button(
                            f"Download all {len(last_batch)} letters in one Word document",
                            combined_bytes, "air_university_internship_letters.docx",
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            type="primary", key=f"bulk_docx_{last_batch[0].get('id')}"
                        )
                        st.download_button(
                            "Download this batch record as Excel", excel_bytes(last_batch), "generated_letter_batch.xlsx",
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"bulk_xlsx_{last_batch[0].get('id')}"
                        )
            except Exception as exc:
                st.error(f"Could not read this file: {exc}")

    # ------------------------------ Search / reprint ------------------------------
    with tabs[2]:
        st.markdown("### Find a previous letter")
        st.caption("Search using the student's exact registration ID. Every matching saved letter can be downloaded again for printing.")
        with st.form("search_form"):
            search_reg = st.text_input("Registration ID", placeholder="Enter registration ID")
            search_button = st.form_submit_button("Search saved letters", type="primary")
        if search_button:
            if not search_reg.strip():
                st.warning("Enter a registration ID first.")
            else:
                try:
                    results = client.table("letters").select("*").eq("registration_id", search_reg.strip()).order("created_at", desc=True).execute().data or []
                    st.session_state["search_results"] = results
                    st.session_state["search_reg"] = search_reg.strip()
                except Exception as exc:
                    st.error(f"Search failed: {exc}")
        results = st.session_state.get("search_results", [])
        if results:
            st.success(f"Found {len(results)} saved letter(s) for registration ID {st.session_state.get('search_reg')}.")
            for record in results:
                with st.container(border=True):
                    left, right = st.columns([4, 1])
                    with left:
                        st.markdown(f"**{record.get('letter_number', '')}** · {record.get('student_name', '')}")
                        st.caption(f"{safe_text(record.get('program'))} · Semester {record.get('semester', '')} · {record.get('company_name') or 'General letter'} · {record.get('letter_date', '')}")
                    with right:
                        st.download_button(
                            "Download / print",
                            build_docx([record]),
                            f"{str(record.get('letter_number', 'letter')).replace('/', '-')}.docx",
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            key=f"reprint_{record.get('id')}",
                            use_container_width=True,
                        )
        elif search_button:
            st.info("No saved letters were found for that registration ID.")

    # ------------------------------ Records / export ------------------------------
    with tabs[3]:
        st.markdown("### Letter register")
        st.write("Download a copy of the saved record data for office filing. The database remains the primary record.")
        if st.button("Load saved records", key="load_records"):
            try:
                with st.spinner("Reading records from Supabase…"):
                    all_records = fetch_all_records(client)
                st.session_state["all_records"] = all_records
            except Exception as exc:
                st.error(f"Could not load records: {exc}")
        all_records = st.session_state.get("all_records", [])
        if all_records:
            records_df = pd.DataFrame(all_records)
            st.metric("Saved letters", len(records_df))
            query = st.text_input("Filter by name, registration ID, or letter number", key="record_filter")
            if query.strip():
                q = query.strip().lower()
                mask = records_df.astype(str).apply(lambda col: col.str.lower().str.contains(re.escape(q), na=False)).any(axis=1)
                view_df = records_df[mask]
            else:
                view_df = records_df
            visible_cols = [c for c in ["letter_number", "letter_date", "student_name", "registration_id", "gender", "program", "semester", "company_name", "city"] if c in view_df.columns]
            st.dataframe(view_df[visible_cols], use_container_width=True, hide_index=True)
            st.download_button(
                "Download all matching records as Excel",
                excel_bytes(view_df.to_dict(orient="records")),
                "air_university_letter_register.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary", key="records_export",
            )
        else:
            st.info("Click **Load saved records** to show and export the Supabase register.")

    st.divider()
    st.caption("Letter wording follows the supplied sample. Verify student details, pronouns, recipient, and letter number before printing or issuing an official letter.")


if __name__ == "__main__":
    app_main()
